"""Excel para la app de precios: exportar la lista, leer un Excel subido y compararlo con los precios actuales."""
import io
import re
import unicodedata

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROJO = "D7141A"
COLORES_ESTADO = {
    "Sube": "FDE2E2", "Baja": "DFF5E6", "Igual": "F1F1F1",
    "No encontrado": "FFF1CC", "Sin precio": "FFF1CC", "Repetido": "FFF1CC",
}


def normalizar(s):
    s = unicodedata.normalize("NFKD", str(s or "").lower()).encode("ascii", "ignore").decode()
    return re.sub(r"[^a-z0-9]+", " ", s).strip()


def a_numero(v):
    """Acepta 1869.36, '1.869,36', '$ 1.869,36', '1869,36'. Devuelve None si no hay número."""
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return None if v != v else float(v)
    s = str(v).strip().replace("$", "").replace(" ", "")
    if not s or s in ("-", "—"):
        return None
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    elif re.fullmatch(r"\d{1,3}(\.\d{3})+", s):
        s = s.replace(".", "")
    try:
        return float(s)
    except ValueError:
        return None


# ---------------------------------------------------------------- exportar
def _hoja(ws, df, monedas=(), anchos=None, resaltar=None, estado_col=None, pct_cols=()):
    borde = Border(bottom=Side(style="thin", color="E8D9BC"))
    ws.append(list(df.columns))
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=ROJO)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 30
    for fila in df.itertuples(index=False):
        ws.append([None if (isinstance(v, float) and v != v) else v for v in fila])
    cols = list(df.columns)
    for j, nombre in enumerate(cols, start=1):
        letra = get_column_letter(j)
        ancho = (anchos or {}).get(nombre) or min(max(len(str(nombre)), *(len(str(v)) for v in df[nombre].head(300)), 8) + 2, 60)
        ws.column_dimensions[letra].width = ancho
        for i in range(2, len(df) + 2):
            celda = ws[f"{letra}{i}"]
            celda.border = borde
            if nombre in monedas:
                celda.number_format = '"$" #,##0.00'
            if nombre in pct_cols:
                celda.number_format = '+0.0"%";-0.0"%";0"%"'
            if resaltar and nombre == resaltar:
                celda.fill = PatternFill("solid", fgColor="FFF6C2")
    if estado_col and estado_col in cols:
        k = cols.index(estado_col)
        for i, fila in enumerate(df.itertuples(index=False), start=2):
            color = COLORES_ESTADO.get(fila[k])
            if color:
                for j in range(1, len(cols) + 1):
                    ws.cell(row=i, column=j).fill = PatternFill("solid", fgColor=color)
    ws.freeze_panes = "B2" if len(cols) > 1 else "A2"
    ws.auto_filter.ref = ws.dimensions


def a_excel(hojas):
    """hojas: lista de (nombre, df, opciones) -> bytes de un .xlsx"""
    wb = Workbook()
    wb.remove(wb.active)
    for nombre, df, op in hojas:
        _hoja(wb.create_sheet(nombre[:31]), df, **(op or {}))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def excel_lista(df_lista, fecha_txt):
    """Lista de precios para descargar. Trae una columna 'Precio nuevo' vacía para completar y volver a subir."""
    df = df_lista.copy()
    df["Precio nuevo"] = None
    instrucciones = pd.DataFrame({"Cómo usar este archivo": [
        f"Lista de precios TAIRA al {fecha_txt}.",
        "Para cambiar precios: escribí el valor en la columna amarilla 'Precio nuevo' y dejá vacías las filas que no cambian.",
        "No borres ni cambies la columna ID TAIRA: se usa para reconocer cada producto.",
        "Después subí el archivo en la app de precios, sección 'Importar Excel', para ver la comparación antes de aprobar.",
    ]})
    return a_excel([
        ("Precios", df, {"monedas": ("Precio actual", "Precio nuevo"), "resaltar": "Precio nuevo", "anchos": {"Producto": 48}}),
        ("Instrucciones", instrucciones, {"anchos": {"Cómo usar este archivo": 110}}),
    ])


# ---------------------------------------------------------------- importar
def leer_archivo(archivo, nombre):
    """Devuelve una tabla con Producto / Medida / Precio (y ID TAIRA, Precio nuevo si están).
    Entiende la lista exportada por la app y también Excel con varias tablas una al lado de la otra."""
    if nombre.lower().endswith(".csv"):
        raw = archivo.read()
        hojas = {}
        for sep in (";", ",", "\t"):
            try:
                df = pd.read_csv(io.BytesIO(raw), sep=sep, encoding="utf-8-sig")
                if df.shape[1] > 1:
                    hojas = {"csv": df}
                    break
            except Exception:
                pass
        if not hojas:
            hojas = {"csv": pd.read_csv(io.BytesIO(raw))}
    else:
        hojas = _leer_xlsx(archivo)
    hojas = {k: _subir_encabezado(v) for k, v in hojas.items()}
    for df in hojas.values():                      # la lista exportada por la app
        if _col(df, ("id taira",)) is not None:
            return df
    bloques = []
    for df in hojas.values():
        cols = list(df.columns)
        inicios = [i for i, c in enumerate(cols) if normalizar(c) in ("producto", "nombre", "descripcion", "articulo")]
        for k, i in enumerate(inicios):
            fin = inicios[k + 1] if k + 1 < len(inicios) else len(cols)
            sub = df.iloc[:, i:fin]
            sub = sub.loc[:, [c for c in sub.columns if normalizar(c) not in ("", "none", "nan")]]
            c_p = _col(sub, ("precio normal", "precio lista", "precio", "importe", "valor"), excluir=("promo", "actual"))
            if c_p is None:
                continue
            t = pd.DataFrame({"Producto": sub.iloc[:, 0], "Precio": sub[c_p]})
            for nom, claves in (("Medida", ("medida",)), ("Ancho", ("ancho",)), ("Largo", ("largo",))):
                c_m = _col(sub, claves)
                if c_m is not None:
                    t[nom] = sub[c_m]
            t = t[t["Producto"].notna() & (t["Producto"].map(normalizar) != "producto") & t["Precio"].map(a_numero).notna()]
            bloques.append(t)
    if bloques:
        return pd.concat(bloques, ignore_index=True)
    return max(hojas.values(), key=len)


def _leer_xlsx(archivo, max_vacias=60):
    """Lee cada hoja cortando cuando aparecen muchas filas vacías seguidas (hay Excel con un millón de filas en blanco)."""
    from openpyxl import load_workbook
    wb = load_workbook(archivo, read_only=True, data_only=True)
    hojas = {}
    for ws in wb.worksheets:
        filas, vacias = [], 0
        for fila in ws.iter_rows(values_only=True):
            if all(v is None or str(v).strip() == "" for v in fila):
                vacias += 1
                if vacias >= max_vacias:
                    break
                continue
            vacias = 0
            filas.append(list(fila))
        if filas:
            ancho = max(len(f) for f in filas)
            filas = [f + [None] * (ancho - len(f)) for f in filas]
            hojas[ws.title] = pd.DataFrame(filas[1:], columns=[str(c) for c in filas[0]])
    wb.close()
    return hojas or {"vacía": pd.DataFrame()}


def _subir_encabezado(df):
    """Si el encabezado real está unas filas más abajo (como en el Excel original), lo detecta."""
    if _col(df, ("producto", "nombre")) is not None:
        return df
    for i in range(min(10, len(df))):
        fila = [normalizar(v) for v in df.iloc[i].tolist()]
        if any(v in ("producto", "nombre", "descripcion", "articulo") for v in fila):
            nuevo = df.iloc[i + 1:].copy()
            nuevo.columns = [str(v) for v in df.iloc[i].tolist()]
            return nuevo.reset_index(drop=True)
    return df


def _col(df, claves, excluir=()):
    for c in df.columns:
        n = normalizar(c)
        if any(n == k or n.startswith(k) for k in claves) and not any(x in n for x in excluir):
            return c
    return None


def comparar(df_in, catalogo):
    """Devuelve (df_comparacion, info). catalogo: lista de dicts con id, nombre, categoria, ancho, largo, precio_publicado."""
    c_id = _col(df_in, ("id taira",))
    c_nom = _col(df_in, ("producto", "nombre", "descripcion", "articulo"))
    c_nuevo = _col(df_in, ("precio nuevo", "nuevo precio", "nuevo"))
    columnas = ["ID", "Producto", "Categoría", "Precio actual", "Precio nuevo", "Diferencia $", "Variación %", "Estado"]
    if c_nuevo is not None and df_in[c_nuevo].map(a_numero).notna().sum() == 0:
        info = {"col_id": c_id, "col_producto": c_nom, "col_precio": c_nuevo}
        return pd.DataFrame(columns=columnas), info   # la columna 'Precio nuevo' está vacía: no hay nada que comparar
    c_precio = c_nuevo or _col(df_in, ("precio normal", "precio lista", "precio", "importe", "valor"), excluir=("actual", "promo"))
    if c_precio is None:
        c_precio = _col(df_in, ("precio",))
    c_meds = [c for c in (_col(df_in, ("medida",)), _col(df_in, ("ancho",)), _col(df_in, ("largo",))) if c is not None]
    info = {"col_id": c_id, "col_producto": c_nom, "col_precio": c_precio}
    if c_precio is None or (c_id is None and c_nom is None):
        return None, info

    por_id = {int(p["id"]): p for p in catalogo}
    por_nombre = {}
    for p in catalogo:
        por_nombre.setdefault(normalizar(p["nombre"]), []).append(p)

    filas, usados = [], {}
    for _, r in df_in.iterrows():
        nombre_x = str(r[c_nom]).strip() if c_nom is not None and pd.notna(r[c_nom]) else ""
        nuevo = a_numero(r[c_precio])
        if not nombre_x and nuevo is None and (c_id is None or pd.isna(r[c_id])):
            continue
        p = None
        if c_id is not None and pd.notna(r[c_id]):
            try:
                p = por_id.get(int(float(r[c_id])))
            except (TypeError, ValueError):
                p = None
        if p is None and nombre_x:
            cands = por_nombre.get(normalizar(nombre_x), [])
            if len(cands) > 1 and c_meds:
                meds = {normalizar(r[c]) for c in c_meds if pd.notna(r[c]) and normalizar(r[c])}
                f = [c for c in cands if meds & {normalizar(c.get("largo")), normalizar(c.get("ancho"))}]
                cands = f if len(f) == 1 else cands
            if len(cands) == 1:
                p = cands[0]
            elif len(cands) > 1:
                filas.append(_fila(None, nombre_x, nuevo, "Repetido"))
                continue
        if p is None:
            if c_nuevo is not None and nuevo is None:
                continue
            filas.append(_fila(None, nombre_x, nuevo, "No encontrado"))
            continue
        if nuevo is None:
            if c_nuevo is not None:
                continue  # fila sin precio nuevo: no cambia
            filas.append(_fila(p, nombre_x, None, "Sin precio"))
            continue
        usados[p["id"]] = usados.get(p["id"], 0) + 1
        actual = p.get("precio_publicado")
        actual = float(actual) if actual is not None else None
        if actual is None:
            estado = "Sube"
        elif abs(nuevo - actual) < 0.011:
            estado = "Igual"
        else:
            estado = "Sube" if nuevo > actual else "Baja"
        filas.append(_fila(p, nombre_x, nuevo, estado, actual))
    df = pd.DataFrame(filas, columns=columnas)
    rep = [i for i, n in usados.items() if n > 1]
    if rep:
        df.loc[df["ID"].isin(rep), "Estado"] = "Repetido"
    info["repetidos_en_excel"] = len(rep)
    return df, info


def _fila(p, nombre_x, nuevo, estado, actual=None):
    nuevo = round(nuevo, 2) if nuevo is not None else None
    dif = (nuevo - actual) if (nuevo is not None and actual is not None) else None
    pct = (dif / actual * 100) if (dif is not None and actual) else None
    return [p["id"] if p else None, p["nombre"] if p else nombre_x, (p.get("categoria") if p else None),
            actual, nuevo, round(dif, 2) if dif is not None else None, round(pct, 1) if pct is not None else None, estado]


def excel_comparacion(df, nombre_archivo, fecha_txt):
    resumen = pd.DataFrame({
        "Dato": ["Archivo comparado", "Fecha", "Productos que suben", "Productos que bajan", "Sin cambio",
                 "No encontrados / a revisar", "Variación promedio (suben y bajan)"],
        "Valor": [nombre_archivo, fecha_txt, int((df["Estado"] == "Sube").sum()), int((df["Estado"] == "Baja").sum()),
                  int((df["Estado"] == "Igual").sum()), int(df["Estado"].isin(["No encontrado", "Repetido", "Sin precio"]).sum()),
                  f"{df.loc[df['Estado'].isin(['Sube', 'Baja']), 'Variación %'].mean():+.1f}%" if df["Estado"].isin(["Sube", "Baja"]).any() else "—"],
    })
    return a_excel([
        ("Comparación", df, {"monedas": ("Precio actual", "Precio nuevo", "Diferencia $"), "pct_cols": ("Variación %",),
                             "estado_col": "Estado", "anchos": {"Producto": 48}}),
        ("Resumen", resumen, {"anchos": {"Dato": 38, "Valor": 40}}),
    ])


def excel_historial(df):
    return a_excel([("Historial", df, {"monedas": ("Precio anterior", "Precio nuevo"), "pct_cols": ("Variación %",),
                                       "anchos": {"Producto": 48}})])
